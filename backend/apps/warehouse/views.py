"""
仓库管理视图
"""
import logging
import io
from django.http import HttpResponse
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from apps.core.exceptions import BusinessException
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
    KitTemplate, KitVersion, KitRequest,
)
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    KitTemplateSerializer, KitVersionSerializer, KitVersionWriteSerializer,
    KitRequestSerializer, KitRequestSubmitSerializer, KitRequestPreviewSerializer,
    StockFreezeSerializer, FreezeInputSerializer, ReturnInputSerializer,
)
from . import services

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 其他视图占位 ====================

class DashboardView(APIView):
    """仪表盘视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'message': '仪表盘功能开发中...'
        })


class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Goods.objects.select_related(
            'variety', 'variety__category'
        ).all().order_by('-created_at')

        keyword = request.query_params.get('keyword')
        if keyword:
            queryset = queryset.filter(name__icontains=keyword)
        variety_id = request.query_params.get('variety')
        if variety_id:
            queryset = queryset.filter(variety_id=variety_id)
        is_active = request.query_params.get('is_active')
        if is_active in ('true', 'false'):
            queryset = queryset.filter(is_active=is_active == 'true')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        goods = queryset[start:end]

        return success_response(data={
            'list': GoodsSerializer(goods, many=True).data,
            'total': total,
            'page': page,
            'page_size': page_size
        })


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StockIn.objects.select_related('goods', 'operator').all().order_by('-stock_in_time')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        records = queryset[start:end]

        return success_response(data={
            'list': StockInSerializer(records, many=True).data,
            'total': total,
            'page': page,
            'page_size': page_size
        })


class StockOutListView(APIView):
    """出库记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StockOut.objects.select_related('goods', 'operator').all().order_by('-created_at')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        records = queryset[start:end]

        return success_response(data={
            'list': StockOutSerializer(records, many=True).data,
            'total': total,
            'page': page,
            'page_size': page_size
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


# ==================== 领用包定义与版本 ====================

def _get_object_or_error(model, pk, message):
    try:
        return model.objects.get(pk=pk)
    except model.DoesNotExist:
        return error_response(message=message, code=404)


def _validate_input(serializer):
    """返回 (validated_data, error_response_or_None)。"""
    if serializer.is_valid():
        return serializer.validated_data, None
    first_error = next(iter(serializer.errors.values()))
    if isinstance(first_error, list):
        first_error = first_error[0]
    if isinstance(first_error, dict):
        first_error = next(iter(first_error.values()))
        if isinstance(first_error, list):
            first_error = first_error[0]
    return None, error_response(message=str(first_error))


class KitTemplateListView(APIView):
    """领用包定义列表/创建"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = KitTemplate.objects.all().order_by('-created_at')
        keyword = request.query_params.get('keyword')
        if keyword:
            queryset = queryset.filter(name__icontains=keyword)
        is_active = request.query_params.get('is_active')
        if is_active in ('true', 'false'):
            queryset = queryset.filter(is_active=is_active == 'true')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        templates = queryset[start:end]
        return success_response(data={
            'list': KitTemplateSerializer(templates, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })

    def post(self, request):
        name = request.data.get('name')
        description = request.data.get('description', '')
        try:
            template = services.create_template(name, description, request.user)
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        logger.info(f"User {request.user.username} created kit template {template.name}")
        return success_response(data=KitTemplateSerializer(template).data, message='创建成功')


class KitTemplateDetailView(APIView):
    """领用包定义详情/更新/删除"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        template = _get_object_or_error(KitTemplate, pk, '领用包不存在')
        if isinstance(template, Response):
            return template
        data = KitTemplateSerializer(template).data
        data['versions'] = KitVersionSerializer(
            template.versions.all(), many=True
        ).data
        return success_response(data=data)

    def put(self, request, pk):
        template = _get_object_or_error(KitTemplate, pk, '领用包不存在')
        if isinstance(template, Response):
            return template
        try:
            template = services.update_template(
                template, request.data.get('name'), request.data.get('description', ''),
                request.user,
            )
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        return success_response(data=KitTemplateSerializer(template).data, message='更新成功')

    def delete(self, request, pk):
        template = _get_object_or_error(KitTemplate, pk, '领用包不存在')
        if isinstance(template, Response):
            return template
        try:
            name = services.delete_template(template, request.user)
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        logger.info(f"User {request.user.username} deleted kit template {name}")
        return success_response(message='删除成功')


class KitVersionView(APIView):
    """领用包版本：创建新版本"""
    permission_classes = [IsAuthenticated]

    def post(self, request, template_pk):
        template = _get_object_or_error(KitTemplate, template_pk, '领用包不存在')
        if isinstance(template, Response):
            return template
        serializer = KitVersionWriteSerializer(data=request.data)
        validated, err = _validate_input(serializer)
        if err:
            return err
        try:
            version = services.create_version(
                template,
                validated.get('items'), validated.get('groups'),
                validated.get('change_remark'), request.user,
                activate=validated.get('activate', False),
            )
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        logger.info(
            f"User {request.user.username} created kit version {version}"
        )
        return success_response(data=KitVersionSerializer(version).data, message='版本创建成功')


class KitVersionDetailView(APIView):
    """版本详情 / 草稿编辑 / 删除 / 生效"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        version = _get_object_or_error(KitVersion, pk, '领用包版本不存在')
        if isinstance(version, Response):
            return version
        return success_response(data=KitVersionSerializer(version).data)

    def put(self, request, pk):
        version = _get_object_or_error(KitVersion, pk, '领用包版本不存在')
        if isinstance(version, Response):
            return version
        serializer = KitVersionWriteSerializer(data=request.data)
        validated, err = _validate_input(serializer)
        if err:
            return err
        try:
            version = services.update_draft_version(
                version, validated.get('change_remark'),
                validated.get('items'), validated.get('groups'), request.user,
            )
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        return success_response(data=KitVersionSerializer(version).data, message='草稿已更新')

    def delete(self, request, pk):
        version = _get_object_or_error(KitVersion, pk, '领用包版本不存在')
        if isinstance(version, Response):
            return version
        try:
            services.delete_version(version, request.user)
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        return success_response(message='删除成功')

    def post(self, request, pk, activate=False):
        """生效版本：POST /kit-versions/<id>/activate/"""
        if not activate:
            return error_response(message='不支持的操作', code=404)
        version = _get_object_or_error(KitVersion, pk, '领用包版本不存在')
        if isinstance(version, Response):
            return version
        try:
            version = services.activate_version(version, request.user)
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        return success_response(data=KitVersionSerializer(version).data, message='版本已生效')


# ==================== 领用申请 ====================

class KitRequestPreviewView(APIView):
    """申请前校验：整套物资是否可交付"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        serializer = KitRequestPreviewSerializer(data=request.data)
        validated, err = _validate_input(serializer)
        if err:
            return err
        try:
            result = services.preview_request(
                validated['version'], validated.get('selections')
            )
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        return success_response(data=result, message=result['message'])


class KitRequestListView(APIView):
    """领用申请列表/提交"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = KitRequest.objects.select_related(
            'template', 'version', 'applicant', 'approved_by'
        ).all().order_by('-created_at')

        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)
        template_id = request.query_params.get('template')
        if template_id:
            queryset = queryset.filter(template_id=template_id)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        records = queryset[start:end]
        return success_response(data={
            'list': KitRequestSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })

    def post(self, request):
        serializer = KitRequestSubmitSerializer(data=request.data)
        validated, err = _validate_input(serializer)
        if err:
            return err
        try:
            kit_request = services.submit_request(
                request.user, validated['version'],
                validated['receiver'], validated.get('receiver_dept'),
                validated.get('remark'), validated.get('selections'),
            )
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        logger.info(
            f"User {request.user.username} submitted kit request {kit_request.number}"
        )
        return success_response(
            data=KitRequestSerializer(kit_request).data, message='申请已提交并占用库存'
        )


class KitRequestDetailView(APIView):
    """申请详情"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        kit_request = _get_object_or_error(KitRequest, pk, '领用申请不存在')
        if isinstance(kit_request, Response):
            return kit_request
        return success_response(data=KitRequestSerializer(kit_request).data)


class KitRequestApproveView(APIView):
    """审批：通过/驳回"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        approved = request.data.get('approved')
        if approved is None:
            return error_response(message='请指定审批结果 approved=true/false')
        opinion = request.data.get('opinion', '')
        try:
            kit_request = services.approve_request(
                pk, request.user, bool(approved), opinion
            )
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        logger.info(
            f"User {request.user.username} approved={approved} kit request {pk}"
        )
        return success_response(
            data=KitRequestSerializer(kit_request).data,
            message='已批准' if approved else '已驳回并释放占用',
        )


class KitRequestCancelView(APIView):
    """撤回待审批申请"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            kit_request = services.cancel_request(pk, request.user)
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        return success_response(
            data=KitRequestSerializer(kit_request).data, message='已撤回并释放占用'
        )


class KitRequestIssueView(APIView):
    """批准后发放（出库）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            kit_request = services.issue_request(pk, request.user)
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        logger.info(f"User {request.user.username} issued kit request {pk}")
        return success_response(
            data=KitRequestSerializer(kit_request).data, message='已发放出库'
        )


class KitRequestReturnView(APIView):
    """部分/全部退回"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        serializer = ReturnInputSerializer(data=request.data)
        validated, err = _validate_input(serializer)
        if err:
            return err
        try:
            kit_request = services.return_request(
                pk, request.user, validated['items']
            )
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        logger.info(f"User {request.user.username} returned kit request {pk}")
        return success_response(
            data=KitRequestSerializer(kit_request).data,
            message='全部退回完成' if kit_request.status == 'returned' else '部分退回已登记',
        )


# ==================== 物资冻结 ====================

class StockFreezeView(APIView):
    """冻结/解冻物资"""
    permission_classes = [IsAuthenticated]

    def post(self, request, freeze_type):
        if freeze_type not in ('freeze', 'unfreeze'):
            return error_response(message='冻结动作非法')
        serializer = FreezeInputSerializer(data=request.data)
        validated, err = _validate_input(serializer)
        if err:
            return err
        try:
            if freeze_type == 'freeze':
                record = services.freeze_goods(
                    validated['goods'], request.user,
                    validated['quantity'], validated['reason'],
                )
                message = '冻结成功'
            else:
                record = services.unfreeze_goods(
                    validated['goods'], request.user,
                    validated['quantity'], validated['reason'],
                )
                message = '解冻成功'
        except BusinessException as exc:
            return error_response(message=str(exc), code=exc.code)
        logger.info(f"User {request.user.username} {freeze_type} goods {validated['goods']}")
        return success_response(data=StockFreezeSerializer(record).data, message=message)


class StockFreezeListView(APIView):
    """冻结记录列表"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        from .models import StockFreeze
        queryset = StockFreeze.objects.select_related('goods', 'operator').all().order_by('-created_at')
        goods_id = request.query_params.get('goods')
        if goods_id:
            queryset = queryset.filter(goods_id=goods_id)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        records = queryset[start:end]
        return success_response(data={
            'list': StockFreezeSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })
